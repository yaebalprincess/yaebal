#!/usr/bin/env python3
"""Build the report and diagrams for Lab 2, variant 6."""
from __future__ import annotations

import csv
import subprocess
import sys
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (Image, PageBreak, Paragraph, Preformatted, SimpleDocTemplate, Spacer, Table,
                                TableStyle)

from lab2_lexer import Lexer, Token, scan_file, write_csv

ROOT = Path(__file__).resolve().parent
FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
FONT_BOLD = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"


def make_dfa() -> None:
    subprocess.run(["dot", "-Tpng", str(ROOT / "dfa.dot"), "-o", str(ROOT / "dfa.png")], check=True)
    subprocess.run(["dot", "-Tsvg", str(ROOT / "dfa.dot"), "-o", str(ROOT / "dfa.svg")], check=True)


def write_demo_outputs() -> None:
    for filename in ("valid_variant6.txt", "errors_variant6.txt"):
        tokens, errors = scan_file(ROOT / filename)
        stem = Path(filename).stem
        write_csv(ROOT / f"{stem}_tokens.csv", tokens)
        (ROOT / f"{stem}_result.txt").write_text(
            f"Файл: {filename}\nРаспознано лексем: {len(tokens)}\nЛексических ошибок: {len(errors)}\n" +
            "\n".join(f"строка {e.line}, столбец {e.column}: {e.message}; фрагмент {e.lexeme!r}" for e in errors) + "\n",
            encoding="utf-8",
        )


def p(text: str, style: ParagraphStyle) -> Paragraph:
    return Paragraph(text, style)


def table_from(rows: list[list[str]], widths: list[float], styles: dict[str, ParagraphStyle], small=False) -> Table:
    body_style = styles["lab_small"] if small else styles["lab_body"]
    data = []
    for row_number, row in enumerate(rows):
        cell_style = styles["lab_head"] if row_number == 0 else body_style
        data.append([p(escape(str(value)), cell_style) for value in row])
    table = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#17334D")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EDF3F7")]),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#B9C8D2")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    return table


def page_decor(canvas, doc):
    canvas.saveState()
    canvas.setStrokeColor(colors.HexColor("#187F89"))
    canvas.setLineWidth(2)
    canvas.line(doc.leftMargin, A4[1] - 28, doc.leftMargin + 24, A4[1] - 28)
    canvas.setFont("DV", 8)
    canvas.setFillColor(colors.HexColor("#566779"))
    canvas.drawString(doc.leftMargin, 18, "Лабораторная работа № 2 · Лексический анализатор · вариант 6")
    canvas.drawRightString(A4[0] - doc.rightMargin, 18, f"{doc.page}")
    canvas.restoreState()


def build_pdf() -> None:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    pdfmetrics.registerFont(TTFont("DV", FONT))
    pdfmetrics.registerFont(TTFont("DVB", FONT_BOLD))
    pdfmetrics.registerFontFamily("DV", normal="DV", bold="DVB")
    styles = getSampleStyleSheet()
    custom_styles = {
        "lab_title": ParagraphStyle("lab_title", parent=styles["Title"], fontName="DVB", fontSize=23, leading=29, textColor=colors.HexColor("#17334D"), spaceAfter=8),
        "lab_h1": ParagraphStyle("lab_h1", parent=styles["Heading1"], fontName="DVB", fontSize=17, leading=21, textColor=colors.HexColor("#17334D"), spaceBefore=8, spaceAfter=8),
        "lab_h2": ParagraphStyle("lab_h2", parent=styles["Heading2"], fontName="DVB", fontSize=12, leading=15, textColor=colors.HexColor("#187F89"), spaceBefore=7, spaceAfter=5),
        "lab_body": ParagraphStyle("lab_body", parent=styles["BodyText"], fontName="DV", fontSize=9, leading=13, textColor=colors.HexColor("#17334D"), spaceAfter=6),
        "lab_small": ParagraphStyle("lab_small", parent=styles["BodyText"], fontName="DV", fontSize=7.6, leading=10, textColor=colors.HexColor("#17334D"), spaceAfter=3),
        "lab_head": ParagraphStyle("lab_head", parent=styles["BodyText"], fontName="DVB", fontSize=7.3, leading=9, textColor=colors.white),
        "lab_code": ParagraphStyle("lab_code", fontName="DV", fontSize=7.2, leading=8.6, textColor=colors.HexColor("#17334D")),
        "lab_center": ParagraphStyle("lab_center", parent=styles["BodyText"], fontName="DV", fontSize=10, leading=14, alignment=TA_CENTER, textColor=colors.HexColor("#566779")),
    }
    for name, style in custom_styles.items():
        styles.add(style)
    doc = SimpleDocTemplate(str(ROOT / "lab2_variant6_report.pdf"), pagesize=A4,
                            rightMargin=1.55 * cm, leftMargin=1.55 * cm, topMargin=1.55 * cm, bottomMargin=1.05 * cm,
                            title="Лабораторная работа № 2 — вариант 6", author="Студент")
    story = []
    story += [Spacer(1, 1.0 * cm), p("Лабораторная работа № 2", styles["lab_title"]),
              p("Проектирование лексического анализатора", styles["lab_h1"]),
              p("Вариант 6 · логические выражения с константами 0 и 1", styles["lab_center"]), Spacer(1, 1.2 * cm)]
    story += [p("Цель работы", styles["lab_h2"]), p("Получить практические навыки построения сканера: выделять лексемы входного текста, распознавать их типы, формировать таблицу лексем и выдавать сообщения о лексических ошибках.", styles["lab_body"])]
    story += [p("1. Задание по лабораторной работе", styles["lab_h1"]),
              p("Входной язык содержит логические выражения, разделённые символом ; (точка с запятой). Логические выражения состоят из идентификаторов, констант 0 и 1, знака присваивания :=, знаков операций or, xor, and, not и круглых скобок.", styles["lab_body"]),
              p("Ограничение: длина идентификатора — не более 32 символов. Комментарии допускаются неограниченной длины; в реализации выбрана форма {…}. Пробельные символы и комментарии удаляются сканером. При обнаружении ошибки сканирование продолжается, чтобы вывести все найденные ошибки.", styles["lab_body"]),
              p("Примечание о приложенных материалах: PDF «2_5 …» содержит учебное условие и варианты, PDF «2_1 Теория» — методические сведения и пример автомата. Lab2.txt и Lab2.exe являются примером реализации, а не отдельным условием варианта 6.", styles["lab_small"]),
              PageBreak()]

    story += [p("2. КС-грамматика входного языка", styles["lab_h1"]),
              p("Синтаксические правила записаны в расширенной форме Бэкуса—Наура. В грамматике задан приоритет операций: not выше and, затем xor, затем or.", styles["lab_body"])]
    grammar = """<Программа> ::= <Выражение> { ; <Выражение> } [ ; ]
<Выражение> ::= <Идентификатор> := <OrExpr> | <OrExpr>
<OrExpr> ::= <XorExpr> { or <XorExpr> }
<XorExpr> ::= <AndExpr> { xor <AndExpr> }
<AndExpr> ::= <UnaryExpr> { and <UnaryExpr> }
<UnaryExpr> ::= not <UnaryExpr> | <Первичный>
<Первичный> ::= <Идентификатор> | 0 | 1 | ( <OrExpr> )

<Идентификатор> ::= <Буква> { <Буква> | <Цифра> | _ }      (длина ≤ 32)
<Буква> ::= A | … | Z | a | … | z
<Цифра> ::= 0 | 1 | … | 9"""
    story += [Preformatted(grammar, styles["lab_code"]), Spacer(1, 5), p("Лексические классы", styles["lab_h2"])]
    lexical_rows = [["Лексема", "Тип", "Условие распознавания"],
                    ["or, xor, and, not", "логическая операция", "точное совпадение строчной записи"],
                    ["0, 1", "логическая константа", "один символ; последовательности 00, 01, 10 и т. п. ошибочны"],
                    [":=", "знак присваивания", "два соседних символа"],
                    [";", "разделитель", "один символ"],
                    ["(, )", "скобка", "один символ"],
                    ["буква…", "идентификатор", "латинская буква, затем буквы/цифры/_, максимум 32"],
                    ["{…}", "комментарий", "неограниченная длина, включая переводы строк; не попадает в таблицу"]]
    story += [table_from(lexical_rows, [4.3 * cm, 4.1 * cm, 8.1 * cm], styles), PageBreak()]

    story += [p("3. Алгоритм и конечный автомат", styles["lab_h1"]),
              p("Сканер выполняет один проход слева направо. Сначала пропускаются пробелы и комментарии. Затем по первому символу выбирается распознаватель: слово, числовая последовательность, присваивание, разделитель, скобка или ошибка. После выделения лексема проверяется, заносится в таблицу, а позиция строки и столбца обновляется. Ошибочный фрагмент поглощается до естественной границы, после чего анализ продолжается.", styles["lab_body"]),
              p("На графе ниже qID и qNUM — допускающие состояния. В qID после распознавания выполняется классификация: точные строки or/xor/and/not становятся операциями, остальные — идентификаторами. Состояние ER означает ошибку. qCOM допускает любой символ до закрывающей скобки комментария.", styles["lab_body"]),
              Image(str(ROOT / "dfa.png"), width=17.2 * cm, height=8.0 * cm),
              p("Рисунок 1 — ДКА выделения лексем варианта 6", styles["lab_small"]), PageBreak()]

    valid_tokens, valid_errors = scan_file(ROOT / "valid_variant6.txt")
    error_tokens, error_errors = scan_file(ROOT / "errors_variant6.txt")
    story += [p("4. Результаты работы программы", styles["lab_h1"]),
              p("Программа запускается командой: python lab2_lexer.py <входной_файл> [--csv table.csv] [--json result.json]. Код возврата равен 0 при отсутствии лексических ошибок, 1 при наличии ошибок и 2 при ошибке чтения файла.", styles["lab_body"]),
              p("4.1. Корректный пример", styles["lab_h2"]),
              Preformatted((ROOT / "valid_variant6.txt").read_text(encoding="utf-8"), styles["lab_code"]),
              p(f"Распознано лексем: {len(valid_tokens)}; лексических ошибок: {len(valid_errors)}.", styles["lab_body"])]
    demo_rows = [["№", "Лексема", "Тип", "Строка", "Столбец"]]
    for token in valid_tokens[:20]:
        demo_rows.append([str(token.number), token.lexeme, token.kind, str(token.line), str(token.column)])
    story += [table_from(demo_rows, [0.8 * cm, 2.4 * cm, 7.7 * cm, 1.5 * cm, 1.7 * cm], styles, small=True),
              p("Показаны первые 20 строк; полная таблица из 34 лексем сохранена в valid_variant6_tokens.csv.", styles["lab_small"]), PageBreak(),
              p("4.2. Проверка ошибок", styles["lab_h2"]),
              Preformatted((ROOT / "errors_variant6.txt").read_text(encoding="utf-8"), styles["lab_code"]),
              p(f"В тесте обнаружено ошибок: {len(error_errors)}. Проверяются недопустимое число 10, идентификатор длиннее 32 символов, недопустимый символ @, идентификатор с начальным _, ошибочное присваивание ':' и незавершённый комментарий.", styles["lab_body"])]
    err_rows = [["Строка", "Столбец", "Сообщение", "Фрагмент"]]
    for error in error_errors:
        err_rows.append([str(error.line), str(error.column), error.message, repr(error.lexeme)])
    story += [table_from(err_rows, [1.5 * cm, 1.6 * cm, 10.0 * cm, 3.0 * cm], styles, small=True), PageBreak()]

    story += [p("5. Текст программы", styles["lab_h1"]),
              p("Ниже приведён полный текст файла lab2_lexer.py. Он также приложен отдельно в репозитории вместе с тестами и примерами входных данных.", styles["lab_body"])]
    code = (ROOT / "lab2_lexer.py").read_text(encoding="utf-8")
    numbered = "\n".join(f"{i:03d}  {line}" for i, line in enumerate(code.splitlines(), 1))
    story.append(Preformatted(numbered, ParagraphStyle("fullcode", fontName="DV", fontSize=6.0, leading=7.2, textColor=colors.HexColor("#17334D"))))
    story += [PageBreak(), p("6. Выводы", styles["lab_h1"]),
              p("В ходе работы разработан лексический анализатор для варианта 6. Он распознаёт идентификаторы, логические константы 0 и 1, операции or/xor/and/not, присваивание :=, разделители и круглые скобки, пропускает комментарии {…}, формирует таблицу лексем с позициями и сообщает об ошибках. Обработка продолжается после ошибки, поэтому пользователь получает полный список проблем во входном файле.", styles["lab_body"]),
              p("Корректность проверена модульными тестами: распознавание всех классов лексем, переход строк после многострочного комментария, восстановление после ошибок и ограничение длины идентификатора. Дополнительно обработан приложенный пример Lab2.txt: сканер сформировал таблицу и обнаружил ошибки, соответствующие ограничениям варианта 6.", styles["lab_body"]),
              p("Исходный код, отчёт, схема автомата, тесты и примеры входных файлов находятся в архиве проекта.", styles["lab_body"])]
    doc.build(story, onFirstPage=page_decor, onLaterPages=page_decor)


def build_docx() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.55)
    section.left_margin = section.right_margin = Inches(0.7)
    normal = doc.styles["Normal"]
    normal.font.name = "DejaVu Sans"
    normal.font.size = Pt(9)
    for name in ("Title", "Heading 1", "Heading 2"):
        doc.styles[name].font.name = "DejaVu Sans"
        doc.styles[name].font.color.rgb = RGBColor(23, 51, 77)
    doc.add_heading("Лабораторная работа № 2", 0)
    doc.add_heading("Проектирование лексического анализатора", 1)
    doc.add_paragraph("Вариант 6 · логические выражения с константами 0 и 1")
    doc.add_paragraph("Входной язык содержит логические выражения, разделённые ';': идентификаторы, 0/1, :=, or/xor/and/not и круглые скобки. Комментарии {…}, длина идентификатора не более 32 символов.")
    doc.add_heading("КС-грамматика", 1)
    doc.add_paragraph("<Программа> ::= <Выражение> { ; <Выражение> } [ ; ]\n<Выражение> ::= <Идентификатор> := <OrExpr> | <OrExpr>\n<OrExpr> ::= <XorExpr> { or <XorExpr> }\n<XorExpr> ::= <AndExpr> { xor <AndExpr> }\n<AndExpr> ::= <UnaryExpr> { and <UnaryExpr> }\n<UnaryExpr> ::= not <UnaryExpr> | <Первичный>\n<Первичный> ::= <Идентификатор> | 0 | 1 | ( <OrExpr> )")
    doc.add_page_break()
    doc.add_heading("Конечный автомат", 1)
    doc.add_picture(str(ROOT / "dfa.png"), width=Inches(6.4))
    doc.add_paragraph("qID и qNUM — допускающие состояния; ключевые слова выделяются классификацией после распознавания слова. ER — состояние ошибки.")
    doc.add_heading("Тесты", 1)
    doc.add_paragraph((ROOT / "valid_variant6.txt").read_text(encoding="utf-8"))
    doc.add_paragraph("Ошибочный тест:")
    doc.add_paragraph((ROOT / "errors_variant6.txt").read_text(encoding="utf-8"))
    doc.add_paragraph("Результаты: корректный пример — 0 ошибок; ошибочный пример — 7 диагностик.")
    doc.add_page_break()
    doc.add_heading("Текст программы", 1)
    code = (ROOT / "lab2_lexer.py").read_text(encoding="utf-8")
    paragraph = doc.add_paragraph()
    for i, line in enumerate(code.splitlines(), 1):
        run = paragraph.add_run(f"{i:03d}  {line}\n")
        run.font.name = "DejaVu Sans Mono"
        run.font.size = Pt(7)
    doc.add_page_break()
    doc.add_heading("Выводы", 1)
    doc.add_paragraph("Разработан и протестирован лексический анализатор варианта 6. Программа строит таблицу лексем с типами, значениями и позициями, пропускает комментарии и сообщает все обнаруженные лексические ошибки.")
    doc.save(ROOT / "lab2_variant6_report.docx")


def main() -> None:
    make_dfa()
    write_demo_outputs()
    build_pdf()
    build_docx()
    print("Report generated: lab2_variant6_report.pdf, lab2_variant6_report.docx")


if __name__ == "__main__":
    main()
