#!/usr/bin/env python3
"""Rebuild the Russian CPM report; requires reportlab, python-docx, PyMuPDF."""
from __future__ import annotations

import csv
import json
import math
from pathlib import Path
from xml.sax.saxutils import escape
from zipfile import ZipFile, ZIP_DEFLATED

import fitz
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph, Table, TableStyle

ROOT = Path(__file__).resolve().parent
FONT_DIR = Path('/usr/share/fonts/truetype/dejavu')
pdfmetrics.registerFont(TTFont('DV', str(FONT_DIR / 'DejaVuSans.ttf')))
pdfmetrics.registerFont(TTFont('DVB', str(FONT_DIR / 'DejaVuSans-Bold.ttf')))
pdfmetrics.registerFontFamily('DV', normal='DV', bold='DVB')
NAVY = '#17334D'
TEAL = '#187F89'
RED = '#C43D46'
GRAY = '#566779'
PALE = '#EDF3F7'

# AOA representation: all original real works occur once. Zero-duration
# dependencies merge their prerequisites without duplicating any real work.
WORKS = [
    ('A1', '1–2', 1, 2, 30, 'Поиск идеи депозитного вклада', 'Иванова И. И.; Петров П. П.', []),
    ('A2', '1–3', 1, 3, 21, 'Разработка системы критериев по выбору идеи', 'Иванов И. И.', []),
    ('A3', '2; 3–4', 3, 4, 2, 'Выбор идеи', 'Иванов И. И.; Петров П. П.', ['A1', 'A2']),
    ('A4', '4–5', 4, 5, 2, 'Технико-экономическое обоснование идеи', 'Петров П. П.', ['A3']),
    ('A5', '5–6', 5, 6, 5, 'Составление проекта договора банка с вкладчиком по депозитному счёту', 'Иванов И. И.', ['A4']),
    ('A6', '6–7', 6, 7, 7, 'Представление проекта договора руководству банка, его рассмотрение и утверждение', 'Иванов И. И.', ['A5']),
    ('A7', '7–8', 7, 8, 20, 'Заказ в типографию на изготовление бланка договора и его выполнение', 'Иванов И. И.', ['A6']),
    ('A8', '7–9', 7, 9, 14, 'Заказ на рекламу в печати и его выполнение', 'Иванов И. И.', ['A6']),
    ('A9', '7–10', 7, 10, 7, 'Заказ на изготовление стенда и его изготовление', 'Иванов И. И.', ['A6']),
    ('A10', '8–11', 8, 11, 1, 'Получение типографского договора', 'Иванов И. И.', ['A7']),
    ('A11', '9; 10–12', 10, 12, 2, 'Оформление зала к приёму клиентов', 'Иванов И. И.', ['A8', 'A9']),
    ('A12', '11; 12–13', 12, 13, 1, 'Семинар с работниками банка', 'Иванов И. И.', ['A10', 'A11']),
    ('A13', '13–14', 13, 14, 0, 'Начало работы с клиентами', 'Не указан', ['A12']),
]
DUMMIES = [('φ1', 2, 3), ('φ2', 9, 10), ('φ3', 11, 12)]
EDGES = [(w[0], w[2], w[3], w[4]) for w in WORKS] + [(k, i, j, 0) for k, i, j in DUMMIES]
EARLY = {i: 0 for i in range(1, 15)}
for j in range(2, 15):
    EARLY[j] = max(EARLY[i] + d for _, i, end, d in EDGES if end == j)
LATE = {14: EARLY[14]}
for i in range(13, 0, -1):
    LATE[i] = min(LATE[j] - d for _, start, j, d in EDGES if start == i)
PARAMS = {}
for key, i, j, duration in EDGES:
    PARAMS[key] = dict(i=i, j=j, duration=duration, ES=EARLY[i], EF=EARLY[i]+duration,
                       LS=LATE[j]-duration, LF=LATE[j],
                       total_float=LATE[j]-EARLY[i]-duration,
                       free_float=EARLY[j]-EARLY[i]-duration)

# Independently validate the AOA transformation using original work predecessors.
AS, AF, ALS, ALF = {}, {}, {}, {}
for key, _, _, _, duration, _, _, predecessors in WORKS:
    AS[key] = max((AF[p] for p in predecessors), default=0)
    AF[key] = AS[key] + duration
for key, _, _, _, duration, _, _, _ in reversed(WORKS):
    successors = [w[0] for w in WORKS if key in w[7]]
    ALF[key] = min((ALS[s] for s in successors), default=AF['A13'])
    ALS[key] = ALF[key] - duration
    p = PARAMS[key]
    assert (p['ES'], p['EF'], p['LS'], p['LF']) == (AS[key], AF[key], ALS[key], ALF[key])
    free = min((AS[s] for s in successors), default=AF['A13']) - AF[key]
    assert p['free_float'] == free
assert EARLY[14] == 68
assert all(p['total_float'] >= p['free_float'] >= 0 for p in PARAMS.values())
assert [w[0] for w in WORKS if PARAMS[w[0]]['total_float'] == 0] == [
    'A1', 'A3', 'A4', 'A5', 'A6', 'A7', 'A10', 'A12', 'A13']

PATHS = []
def visit(node, nodes, length):
    if node == 14:
        PATHS.append((nodes, length))
    for _, i, j, d in EDGES:
        if i == node:
            visit(j, nodes+[j], length+d)
visit(1, [1], 0)
PATHS.sort(key=lambda p: -p[1])
assert sorted(length for _, length in PATHS) == [47, 54, 56, 59, 63, 68]

class Drawing:
    """One geometry, two vector outputs (PDF and editable SVG)."""
    width, height = 1120, 640
    def __init__(self):
        self.c = canvas.Canvas(str(ROOT/'network-diagram.pdf'), pagesize=(self.width, self.height))
        self.c.setTitle('Сетевой график проекта — 68 дней')
        self.svg = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{self.width}" height="{self.height}" viewBox="0 0 {self.width} {self.height}">',
                    '<rect width="100%" height="100%" fill="white"/>']
    def text(self, x, y, text, size=15, color=NAVY, bold=False, anchor='middle', background=False):
        font = 'DVB' if bold else 'DV'
        width = pdfmetrics.stringWidth(text, font, size)
        if background:
            left = x-width/2-5 if anchor == 'middle' else x-5
            self.c.setFillColor(colors.white)
            self.c.rect(left, self.height-y-4, width+10, size+7, stroke=0, fill=1)
            self.svg.append(f'<rect x="{left}" y="{y-size-3}" width="{width+10}" height="{size+7}" fill="white"/>')
        self.c.setFillColor(colors.HexColor(color))
        self.c.setFont(font, size)
        (self.c.drawCentredString if anchor == 'middle' else self.c.drawString)(x, self.height-y, text)
        self.svg.append(f'<text x="{x}" y="{y}" text-anchor="{anchor}" font-family="DejaVu Sans, sans-serif" font-size="{size}" font-weight="{700 if bold else 400}" fill="{color}">{escape(text)}</text>')
    def line(self, points, color=GRAY, dashed=False, arrow=True, curve=False, width=2.5):
        self.c.setStrokeColor(colors.HexColor(color))
        self.c.setFillColor(colors.HexColor(color))
        self.c.setLineWidth(width)
        self.c.setDash([8, 6] if dashed else [])
        p=self.c.beginPath()
        p.moveTo(points[0][0], self.height-points[0][1])
        if curve:
            p.curveTo(*[v for x,y in points[1:] for v in (x, self.height-y)])
            d=f'M {points[0][0]} {points[0][1]} C '+' '.join(f'{x} {y}' for x,y in points[1:])
        else:
            for x,y in points[1:]: p.lineTo(x, self.height-y)
            d='M '+' L '.join(f'{x} {y}' for x,y in points)
        self.c.drawPath(p)
        dash=' stroke-dasharray="8 6"' if dashed else ''
        self.svg.append(f'<path d="{d}" fill="none" stroke="{color}" stroke-width="{width}"{dash}/>')
        self.c.setDash([])
        if arrow:
            x,y=points[-1]; px,py=points[-2]
            norm=math.hypot(x-px,y-py); ux=(x-px)/norm; uy=(y-py)/norm
            tri=[(x,y),(x-12*ux+5*uy,y-12*uy-5*ux),(x-12*ux-5*uy,y-12*uy+5*ux)]
            p=self.c.beginPath(); p.moveTo(tri[0][0],self.height-tri[0][1])
            for tx,ty in tri[1:]: p.lineTo(tx,self.height-ty)
            p.close(); self.c.drawPath(p,stroke=0,fill=1)
            self.svg.append(f'<polygon points="'+ ' '.join(f'{tx},{ty}' for tx,ty in tri)+f'" fill="{color}"/>')
    def node(self, key, x, y):
        col=RED if EARLY[key] == LATE[key] else TEAL
        self.c.setFillColor(colors.white); self.c.setStrokeColor(colors.HexColor(col)); self.c.setLineWidth(2.5)
        self.c.circle(x,self.height-y,28,fill=1,stroke=1)
        self.svg.append(f'<circle cx="{x}" cy="{y}" r="28" fill="white" stroke="{col}" stroke-width="2.5"/>')
        self.line([(x-27,y+1),(x+27,y+1)],col,arrow=False,width=1)
        self.line([(x,y+1),(x,y+27)],col,arrow=False,width=1)
        self.text(x,y-7,str(key),17,col,bold=True)
        self.text(x-13,y+19,str(EARLY[key]),14,col)
        self.text(x+13,y+19,str(LATE[key]),14,col)
    def finish(self):
        self.c.showPage(); self.c.save()
        self.svg.append('</svg>')
        (ROOT/'network-diagram.svg').write_text('\n'.join(self.svg),encoding='utf-8')

d=Drawing()
d.text(40,25,'СЕТЕВОЙ ГРАФИК · НОВЫЙ ДЕПОЗИТНЫЙ СЧЁТ',18,bold=True,anchor='start')
positions={1:(70,145),2:(220,145),3:(370,145),4:(520,145),5:(670,145),6:(820,145),7:(970,145),
           8:(970,320),9:(785,320),10:(600,320),11:(970,500),12:(600,500),13:(370,500),14:(170,500)}
geometry={
 'A1': ([(98,145),(192,145)],(145,126),False),
 'A2': ([(70,117),(70,38),(370,38),(370,117)],(220,60),True),
 'φ1': ([(248,145),(342,145)],(295,126),False),
 'A3': ([(398,145),(492,145)],(445,126),False),
 'A4': ([(548,145),(642,145)],(595,126),False),
 'A5': ([(698,145),(792,145)],(745,126),False),
 'A6': ([(848,145),(942,145)],(895,126),False),
 'A7': ([(970,173),(970,292)],(1030,242),False),
 'A8': ([(949,164),(805,301)],(913,261),False),
 'A9': ([(943,152),(830,225),(600,200),(600,292)],(717,211),True),
 'φ2': ([(757,320),(628,320)],(693,302),False),
 'A10': ([(970,348),(970,472)],(1031,417),False),
 'A11': ([(600,348),(600,472)],(664,417),False),
 'φ3': ([(942,500),(628,500)],(785,481),False),
 'A12': ([(572,500),(398,500)],(485,481),False),
 'A13': ([(342,500),(198,500)],(270,481),False),
}
for key,i,j,duration in EDGES:
    points,label,curve=geometry[key]
    critical=PARAMS[key]['total_float']==0
    col=RED if critical else TEAL
    d.line(points,col,key.startswith('φ'),curve=curve,width=3 if critical else 2.3)
    d.text(*label,f'{key} · {duration} дн.',14,col,background=True)
for n,(x,y) in positions.items(): d.node(n,x,y)
d.text(65,292,'Критический путь — 68 дней',22,RED,True,anchor='start')
d.text(65,330,'1 → 2 → 3 → 4 → 5 → 6 → 7 → 8 →',16,anchor='start')
d.text(65,357,'11 → 12 → 13 → 14',16,anchor='start')
d.text(65,397,'В круге: номер события / ранний | поздний срок.',14,GRAY,anchor='start')
d.text(65,422,'Время отсчитывается от события 1; Tр(1) = 0.',14,GRAY,anchor='start')
d.line([(65,570),(110,570)],RED)
d.text(124,575,'Критические работы и связи',14,anchor='start')
d.line([(550,570),(595,570)],TEAL,dashed=True)
d.text(609,575,'Фиктивные связи φ1–φ3: 0 дней',14,anchor='start')
d.text(65,611,'A1–A13 — работы из исходной таблицы; красный цвет означает нулевой полный резерв.',14,GRAY,anchor='start')
d.finish()
with fitz.open(ROOT/'network-diagram.pdf') as diagram:
    diagram[0].get_pixmap(matrix=fitz.Matrix(2,2),alpha=False).save(ROOT/'network-diagram.png')

for filename,rows in [('activity-parameters.csv', [dict(code=k,**PARAMS[k]) for k,_,_,_ in EDGES]),
                      ('event-parameters.csv',[dict(event=i,early=EARLY[i],late=LATE[i],reserve=LATE[i]-EARLY[i]) for i in range(1,15)])]:
    with (ROOT/filename).open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=list(rows[0])); writer.writeheader(); writer.writerows(rows)
(ROOT/'calculations.json').write_text(json.dumps(dict(duration=68,events={i:dict(early=EARLY[i],late=LATE[i]) for i in EARLY},
    activities=PARAMS,paths=[dict(events=p,duration=t) for p,t in PATHS]),ensure_ascii=False,indent=2),encoding='utf-8')

# Editable Graphviz source includes the same node times and edge durations.
dot=['digraph project {','  rankdir=LR;','  graph [bgcolor="white", label="Сетевой график: 68 дней", labelloc=t];',
     '  node [shape=circle, fontname="DejaVu Sans"];','  edge [fontname="DejaVu Sans"];']
for i in range(1,15):
    dot.append(f'  {i} [label="{i}\\n{EARLY[i]} | {LATE[i]}"];')
for key,i,j,duration in EDGES:
    col=RED if PARAMS[key]['total_float']==0 else TEAL
    style='dashed' if key.startswith('φ') else 'solid'
    dot.append(f'  {i} -> {j} [label="{key}: {duration}", color="{col}", style={style}];')
dot.append('}')
(ROOT/'network-diagram.dot').write_text('\n'.join(dot)+'\n',encoding='utf-8')

SOURCE_NOTE='Источник: Zadanie_3.pdf, стр. 1; внутри файла заголовок «Задание 4 (Построение сетевого графика)».'
ASSUMPTIONS=[
 'За начало отсчёта принято завершение разработки плана мероприятий: событие 1, время 0. Длительность разработки плана в исходнике не задана, поэтому она не включена в 68 дней.',
 'Зависимости «2; 3», «9; 10» и «11; 12» означают ожидание обеих работ. Сроки выражены в относительных днях; календарная дата начала и режим выходных не заданы.',
 'Применён метод критического пути (CPM). Параллельное выполнение разрешено; ограничения занятости исполнителей в исходных данных не формализованы.'
]
SOURCE_ROWS=[['Код','В исходнике i–j','Содержание работы','Дни','Исполнитель'],
             ['—','— → 1','Разработка плана мероприятий','—','Иванова И. И.']]
SOURCE_ROWS += [[w[0],w[1],w[5],str(w[4]),w[6]] for w in WORKS]
EVENT_ROWS=[['Событие','Ранний срок Tр','Поздний срок Tп','Резерв R','Критическое']]
EVENT_ROWS += [[str(i),str(EARLY[i]),str(LATE[i]),str(LATE[i]-EARLY[i]),'Да' if EARLY[i]==LATE[i] else 'Нет'] for i in range(1,15)]
ACT_ROWS=[['Работа','Дуга i–j','t','РН (ES)','РО (EF)','ПН (LS)','ПО (LF)','Rп','Rс']]
for key,i,j,duration in EDGES:
    p=PARAMS[key]
    ACT_ROWS.append([key,f'{i}–{j}',str(duration),*[str(p[k]) for k in ['ES','EF','LS','LF','total_float','free_float']]])
PATH_ROWS=[['№','Полный путь (события)','Длительность','Резерв пути']]
PATH_ROWS += [[str(k),' → '.join(map(str,p)),str(t),str(68-t)] for k,(p,t) in enumerate(PATHS,1)]
MERGES=[
 'φ1: 2 → 3, 0 дней. A3 (выбор идеи) выполняется по дуге 3 → 4 только после A1 и A2.',
 'φ2: 9 → 10, 0 дней. A11 (оформление зала) выполняется по дуге 10 → 12 после A8 и A9.',
 'φ3: 11 → 12, 0 дней. A12 (семинар) выполняется по дуге 12 → 13 после A10 и A11.',
 'После объединения событие 3 означает готовность к выбору идеи, 10 — готовность к оформлению зала, 12 — готовность к семинару. Окончание отдельной входящей работы может предшествовать сроку общего события; это отражено свободным резервом.'
]
CONCLUSIONS=[
 'Минимальная длительность проекта при заданных зависимостях — 68 дней от события 1. Начало работы с клиентами (A13) — в момент 68; его длительность равна нулю.',
 'Единственный критический путь: 1 → 2 → 3 → 4 → 5 → 6 → 7 → 8 → 11 → 12 → 13 → 14. Фиктивные связи φ1 и φ3 входят в этот путь и не увеличивают его длину.',
 'Критические реальные работы: A1, A3, A4, A5, A6, A7, A10, A12 и A13. Проверка суммы: 30 + 2 + 2 + 5 + 7 + 20 + 1 + 1 + 0 = 68 дней. Задержка любой из них при неизменном остальном плане сдвигает окончание проекта.',
 'Полные резервы некритических работ: A2 — 9 дней; A8 — 5 дней; A9 — 12 дней; A11 — 5 дней. Резервы последовательных работ нельзя складывать: например, A8 и A11 используют один запас в 5 дней на своей ветви.',
 'Свободные резервы: A2 — 9 дней, A9 — 7 дней, A11 — 5 дней; у остальных реальных работ — 0. A8 можно сдвинуть до 5 дней без задержки проекта, но это сдвинет раннее начало следующей работы A11.',
 'События 9 и 10 имеют резерв по 5 дней; остальные события — нулевой. Ближайший к критическому полный путь длится 63 дня, его запас — 5 дней.',
 'Приоритет контроля: поиск идеи, согласование договора и изготовление бланков. Если требуется сократить срок проекта, сокращать нужно работы критического пути, затем пересчитывать сеть.',
 'Ограничение результата: в таблице несколько параллельных работ закреплены за Ивановым И. И. Их длительности включают заказы и ожидание, но трудозатраты и доступность исполнителя не заданы. Поэтому 68 дней — срок по логическим зависимостям; ресурсное выравнивание требует дополнительных данных.'
]

# Five printable landscape A4 pages, including a vector network diagram.
PW,PH=landscape(A4); M=32; CW=PW-2*M
styles={
 'body':ParagraphStyle('body',fontName='DV',fontSize=9,leading=13,textColor=colors.HexColor(NAVY)),
 'small':ParagraphStyle('small',fontName='DV',fontSize=8,leading=11,textColor=colors.HexColor(GRAY)),
 'cell':ParagraphStyle('cell',fontName='DV',fontSize=7.5,leading=10,textColor=colors.HexColor(NAVY)),
 'head':ParagraphStyle('head',fontName='DVB',fontSize=8,leading=11,textColor=colors.white),
}
c=canvas.Canvas(str(ROOT/'project-network-report.pdf'),pagesize=(PW,PH),pageCompression=1)
c.setTitle('Сетевой график проекта и анализ — новый банковский депозитный счёт')
c.setAuthor('Расчёт по исходным данным Zadanie_3.pdf')
def paragraph(text,y,small=False):
    p=Paragraph(escape(text),styles['small' if small else 'body']); _,h=p.wrap(CW,PH)
    p.drawOn(c,M,y-h); return y-h-7
def page(title,subtitle,num):
    c.setFillColor(colors.HexColor(TEAL));c.rect(M,PH-34,32,4,stroke=0,fill=1)
    c.setFont('DVB',21);c.setFillColor(colors.HexColor(NAVY));c.drawString(M,PH-65,title)
    c.setFont('DV',9);c.setFillColor(colors.HexColor(GRAY));c.drawString(M,PH-84,subtitle)
    c.setStrokeColor(colors.HexColor('#D7E2E9'));c.line(M,29,PW-M,29)
    c.setFont('DV',8);c.drawString(M,16,'Новый банковский депозитный счёт · метод критического пути (CPM)')
    c.drawRightString(PW-M,16,f'{num} / 5')
    return PH-102
def table(rows,widths,y,font_size=None):
    data=[[Paragraph(escape(str(v)),styles['head' if k==0 else 'cell']) for v in row] for k,row in enumerate(rows)]
    t=Table(data,colWidths=widths,repeatRows=1)
    t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor(NAVY)),('VALIGN',(0,0),(-1,-1),'MIDDLE'),
        ('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor(PALE)]),
        ('LEFTPADDING',(0,0),(-1,-1),7),('RIGHTPADDING',(0,0),(-1,-1),7),
        ('TOPPADDING',(0,0),(-1,-1),4),('BOTTOMPADDING',(0,0),(-1,-1),4),
        ('LINEBELOW',(0,0),(-1,0),.5,colors.HexColor(NAVY))]))
    _,h=t.wrap(CW,PH)
    assert y-h>37, f'Table exceeds page: {y-h}'
    t.drawOn(c,M,y-h); return y-h-10

y=page('Сетевой график проекта','Выпуск нового банковского депозитного счёта · исходные данные и принятые условия',1)
y=paragraph(SOURCE_NOTE,y,True)
y=table(SOURCE_ROWS,[39,76,CW-39-76-37-155,37,155],y)
for text in ASSUMPTIONS: y=paragraph(text,y,True)
assert y>32,y
c.showPage()
y=page('Сетевая модель проекта','Работы — стрелки; события — круги; длительности — в днях. Красным выделен критический путь.',2)
# Insert the vector diagram after the report has been generated.
diagram_rect=fitz.Rect(M,110,PW-M,PH-54)
c.showPage()
y=page('Сроки наступления событий','Прямой и обратный проходы по сети; начало отсчёта Tр(1) = 0.',3)
y=table(EVENT_ROWS,[85,170,170,150,CW-575],y)
for text in [
 'Прямой проход: Tр(j) = max [Tр(i) + t(i,j)] по всем входящим дугам. Обратный проход: Tп(i) = min [Tп(j) − t(i,j)] по всем исходящим дугам; Tп(14) = 68. Резерв события: R(i) = Tп(i) − Tр(i).',
 'Примеры: Tр(3) = max(21; 30 + 0) = 30; Tр(10) = max(46 + 7; 60 + 0) = 60; Tр(12) = max(60 + 2; 67 + 0) = 67.',
 'Проверка обратного прохода: Tп(7) = min(66 − 20; 65 − 14; 65 − 7) = 46. У событий 9 и 10 резерв 5 дней; у остальных — 0.'
]: y=paragraph(text,y,True)
for text in MERGES: y=paragraph(text,y,True)
assert y>32,y
c.showPage()
y=page('Параметры работ и резервы времени','РН/РО — ранние начало/окончание; ПН/ПО — поздние начало/окончание; Rп/Rс — полный/свободный резерв.',4)
y=table(ACT_ROWS,[65,80,45]+[(CW-190)/6]*6,y)
for text in [
 'РН(i,j) = Tр(i); РО = Tр(i) + t; ПН = Tп(j) − t; ПО = Tп(j). Полный резерв Rп = Tп(j) − Tр(i) − t; свободный резерв Rс = Tр(j) − Tр(i) − t.',
 'Пример A9: РН = 46; РО = 53; ПН = 65 − 7 = 58; ПО = 65; Rп = 65 − 46 − 7 = 12; Rс = 60 − 46 − 7 = 7 дней.',
 'φ1–φ3 — фиктивные связи; они не требуют времени или ресурсов. A13 — реальная завершающая веха с заданной в условии нулевой длительностью.'
]: y=paragraph(text,y,True)
assert y>32,y
c.showPage()
y=page('Полные пути и выводы','Полный путь соединяет исходное событие 1 с завершающим событием 14. Всего в сети 6 полных путей.',5)
y=table(PATH_ROWS,[30,CW-215,90,95],y)
for text in CONCLUSIONS: y=paragraph(text,y,True)
assert y>32,y
c.showPage();c.save()
with fitz.open(ROOT/'project-network-report.pdf') as report, fitz.open(ROOT/'network-diagram.pdf') as diagram:
    report[1].show_pdf_page(diagram_rect,diagram,0)
    report.save(ROOT/'project-network-report.tmp.pdf',garbage=4,deflate=True)
(ROOT/'project-network-report.tmp.pdf').replace(ROOT/'project-network-report.pdf')

# Editable document, containing the same numerical results and an embedded diagram.
doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(11.69); sec.page_height=Inches(8.27)
sec.top_margin=sec.bottom_margin=Inches(.48)
sec.left_margin=sec.right_margin=Inches(.5)
normal=doc.styles['Normal'];normal.font.name='DejaVu Sans';normal.font.size=Pt(9)
normal.paragraph_format.space_after=Pt(5)
for name in ['Heading 1','Heading 2']:
    doc.styles[name].font.name='DejaVu Sans';doc.styles[name].font.color.rgb=RGBColor.from_string('17334D')
def heading(title,sub):
    doc.add_heading(title,0);doc.add_paragraph(sub)
def doc_table(rows,size=8):
    t=doc.add_table(rows=0,cols=len(rows[0]));t.style='Light Shading Accent 1'
    for row in rows:
        cells=t.add_row().cells
        for cell,text in zip(cells,row):
            cell.text=str(text)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(1)
                for r in p.runs:r.font.size=Pt(size)
    return t
heading('Сетевой график проекта','Выпуск нового банковского депозитного счёта. Расчёт и анализ методом CPM.')
doc.add_paragraph(SOURCE_NOTE)
doc_table(SOURCE_ROWS,7)
for text in ASSUMPTIONS:doc.add_paragraph(text)
doc.add_page_break()
heading('Сетевой график и логические связи','Длительность проекта — 68 дней. Красным выделен критический путь.')
doc.add_picture(str(ROOT/'network-diagram.png'),width=Inches(10.2))
doc.add_page_break()
heading('Сроки событий','Tр — ранний срок, Tп — поздний срок; R = Tп − Tр.')
doc_table(EVENT_ROWS)
doc.add_paragraph('Прямой проход: Tр(j) = max[Tр(i) + t(i,j)]. Обратный проход: Tп(i) = min[Tп(j) − t(i,j)], Tп(14) = 68.')
doc.add_paragraph('Tр(3) = max(21; 30) = 30; Tр(10) = max(53; 60) = 60; Tр(12) = max(62; 67) = 67.')
doc.add_page_break()
heading('Параметры работ','Все значения приведены в днях; φ1–φ3 — фиктивные связи.')
doc_table(ACT_ROWS)
doc.add_paragraph('РН = Tр(i); РО = РН + t; ПН = Tп(j) − t; ПО = Tп(j). Rп = Tп(j) − Tр(i) − t; Rс = Tр(j) − Tр(i) − t.')
doc.add_paragraph('Для A9: РН = 46; РО = 53; ПН = 58; ПО = 65; Rп = 12; Rс = 7 дней.')
doc.add_page_break()
heading('Полные пути и выводы','Критический путь имеет наибольшую длительность и нулевой резерв.')
doc_table(PATH_ROWS)
for text in CONCLUSIONS:doc.add_paragraph(text)
doc.add_page_break()
heading('Пояснение построения сети','Сохранена нумерация событий 1–14. Реальные работы не дублируются.')
for text in MERGES:doc.add_paragraph(text)
doc.add_paragraph('Дуги, уточняющие строки исходной таблицы: «2; 3 → 4» представлена работой A3: 3 → 4 и связью φ1; «9; 10 → 12» — A11: 10 → 12 и φ2; «11; 12 → 13» — A12: 12 → 13 и φ3.')
doc.add_paragraph('Эта запись эквивалентна исходным зависимостям: выбор идеи ждёт поиск и критерии, оформление зала — рекламу и стенд, семинар — получение договора и оформление зала. События объединения не следует трактовать как отдельные ранние окончания всех входящих работ.')
doc.core_properties.title='Сетевой график проекта и анализ'
doc.core_properties.subject='Новый банковский депозитный счёт; срок проекта 68 дней'
doc.save(ROOT/'project-network-report.docx')

README='''# Сетевой график проекта

Выполнено задание из **Zadanie_3.pdf** (на странице заголовок «Задание 4»): построен сетевой график выпуска нового банковского депозитного счёта и проведён анализ методом критического пути.

**Срок проекта — 68 дней** от завершения разработки плана мероприятий. Длительность самого планирования не задана. Критические работы: A1 → A3 → A4 → A5 → A6 → A7 → A10 → A12 → A13.

## Готовые файлы

- [Отчёт PDF](project-network-report.pdf) — исходные данные, график, сроки событий, параметры работ, все полные пути и выводы.
- [Редактируемый отчёт DOCX](project-network-report.docx).
- [График SVG](network-diagram.svg) — масштабируется и редактируется в графических редакторах.
- [График PDF](network-diagram.pdf) и [PNG](network-diagram.png).
- [Архив всех материалов](network-project.zip).
- CSV с [параметрами работ](activity-parameters.csv) и [сроками событий](event-parameters.csv), [расчёты JSON](calculations.json), [исходник Graphviz](network-diagram.dot).

## Как читать график

В круге сверху указан номер события, снизу слева — ранний срок, справа — поздний. У стрелки указаны код работы и длительность. Красным выделены критические работы и связи. Пунктир — фиктивные связи нулевой длительности.

Для сохранения нумерации исходных событий введены три зависимости:

- φ1: 2 → 3, после которой выбор идеи A3 идёт по дуге 3 → 4;
- φ2: 9 → 10, после которой оформление зала A11 идёт по дуге 10 → 12;
- φ3: 11 → 12, после которой семинар A12 идёт по дуге 12 → 13.

События 3, 10 и 12 обозначают объединённую готовность соответствующих ветвей. Входящие реальные работы могут закончиться раньше наступления общего события. Это корректно учитывается свободным резервом и не добавляет зависимостей между реальными работами.

Полные резервы: A2 — 9 дней, A8 — 5, A9 — 12, A11 — 5. Запасы последовательных работ нельзя складывать. Ограничения занятости исполнителей не заданы; 68 дней — результат по логическим зависимостям без ресурсного выравнивания. Материалы лекций отдельно не предоставлены; использованы стандартные формулы CPM.

## Воспроизведение

Нужны Python 3, ReportLab 4.4.9, python-docx 1.2.0, PyMuPDF 1.26.6 и системный шрифт DejaVu Sans. Установка в текущей среде не потребовалась — зависимости уже доступны.

```bash
cd /workspace/yaebal/network-project
python build_report.py
```

Скрипт пересчитывает сеть прямым и обратным проходами, независимо сверяет каждую реальную работу с моделью её исходных предшественников и проверяет все шесть полных путей. Работает без сетевых запросов.
'''
(ROOT/'README.md').write_text(README,encoding='utf-8')
files=['project-network-report.pdf','project-network-report.docx','network-diagram.pdf','network-diagram.svg','network-diagram.png',
       'activity-parameters.csv','event-parameters.csv','calculations.json','network-diagram.dot','build_report.py','README.md']
with ZipFile(ROOT/'network-project.zip','w',ZIP_DEFLATED) as archive:
    for filename in files:archive.write(ROOT/filename,arcname='network-project/'+filename)
print('CPM verified: 68 days; 13 real works; 3 dummy links; 14 events; 6 full paths.')
print('Independent predecessor model: all early/late times and free reserves agree.')
for filename in files+['network-project.zip']:print(filename,(ROOT/filename).stat().st_size)
